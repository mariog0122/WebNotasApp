"""Convierte las plantillas Word institucionales (sintaxis Jinja/docxtpl) a sintaxis docxtemplater.

Uso: python3 convert_templates.py <plantilla_mapeada.docx> <plantilla_micro.docx> <carpeta_salida>
(requiere: pip install python-docx; salida usada: public/plantillas/pca.docx y public/plantillas/micro.docx)

Solo cambia el texto de los marcadores; formato, logos, encabezados y pies de página se conservan.
- {% for u in X %} ... {{ u.campo }} ... {% endfor %}  ->  {{#X}} ... {{campo}} ... {{/X}}
- Años fijos del título ("2024 – 2025")                 ->  {{ANIO_LECTIVO_TITULO}}
- Filas de ejemplo vacías                               ->  una sola fila repetible por bucle
"""
import re
import sys

import docx
from docx.oxml.ns import qn

YEAR_RX = re.compile(r'\d{4}\s*[–-]\s*\d{4}')


def paragraphs_of(document):
    for table in document.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from cell.paragraphs
    yield from document.paragraphs


def set_text(paragraph, text):
    """Deja todo el texto en el primer run con texto, sin tocar runs con imágenes."""
    runs = [r for r in paragraph.runs if r._r.find(qn('w:t')) is not None]
    if not runs:
        return
    runs[0].text = text
    for run in runs[1:]:
        for t in run._r.findall(qn('w:t')):
            t.text = ''


def convert_jinja(document):
    loop_stack = []
    for p in paragraphs_of(document):
        text = p.text
        if '{%' not in text and not re.search(r'\{\{\s*\w+\.\w+\s*\}\}', text):
            continue

        def open_loop(m):
            loop_stack.append(m.group(2))
            return '{{#' + m.group(2) + '}}'

        new = re.sub(r'\{%\s*for\s+(\w+)\s+in\s+(\w+)\s*%\}\s*', open_loop, text)
        new = re.sub(r'\{\{\s*\w+\.(\w+)\s*\}\}', r'{{\1}}', new)
        new = re.sub(r'\{%\s*endfor\s*%\}', lambda m: '{{/' + loop_stack.pop() + '}}', new)
        set_text(p, new.strip())


def replace_in_paragraphs(document, pattern, repl):
    for p in paragraphs_of(document):
        if re.search(pattern, p.text):
            set_text(p, re.sub(pattern, repl, p.text))


def table_rows_with(table, needle):
    return [row for row in table.rows if needle in ''.join(c.text for c in row.cells)]


def make_row_loop(table, first_tag, last_tag, loop):
    row = table_rows_with(table, first_tag)[0]
    cells = []
    for c in row.cells:
        if not cells or c._tc is not cells[-1]._tc:
            cells.append(c)
    first_p = next(p for p in cells[0].paragraphs if first_tag in p.text)
    set_text(first_p, first_p.text.replace(first_tag, '{{#' + loop + '}}' + first_tag))
    last_p = next(p for p in cells[-1].paragraphs if last_tag in p.text)
    set_text(last_p, last_p.text.replace(last_tag, last_tag + '{{/' + loop + '}}'))
    # Quita las filas vacías de ejemplo que siguen a la fila del bucle.
    tr = row._tr
    nxt = tr.getnext()
    while nxt is not None and nxt.tag == qn('w:tr') and not ''.join(t.text or '' for t in nxt.iter(qn('w:t'))).strip():
        after = nxt.getnext()
        tr.getparent().remove(nxt)
        nxt = after


def convert_pca(src, dst):
    d = docx.Document(src)
    convert_jinja(d)
    replace_in_paragraphs(d, r'PLANIFICACIÓN CURRICULAR ANUAL\s*' + YEAR_RX.pattern, 'PLANIFICACIÓN CURRICULAR ANUAL {{ANIO_LECTIVO_TITULO}}')
    table = d.tables[0]
    # Tres filas repetían {{OBJETIVOS_SUBNIVEL}}: se deja una sola con los objetivos en líneas.
    rows = table_rows_with(table, '{{OBJETIVOS_SUBNIVEL}}')
    for row in rows[1:]:
        table._tbl.remove(row._tr)
    d.save(dst)


def convert_micro(src, dst):
    d = docx.Document(src)
    convert_jinja(d)
    replace_in_paragraphs(d, r'PLANIFICACIÓN MICROCURRICULAR\s*' + YEAR_RX.pattern, 'PLANIFICACIÓN MICROCURRICULAR {{ANIO_LECTIVO_TITULO}}')
    replace_in_paragraphs(d, r'^\s*PRIMER TRIMESTRE\s*$', '{{TRIMESTRE}}')
    replace_in_paragraphs(d, r'^\s*LOGO\s*$', '')
    table = d.tables[0]
    make_row_loop(table, '{{SEMANA}}', '{{ACTIVIDADES}}', 'SEMANAS')
    make_row_loop(table, '{{INSERCION_TIPO}}', '{{INSERCION_INDICADORES}}', 'INSERCIONES')
    make_row_loop(table, '{{NEE_CONTENIDOS}}', '{{NEE_ACTIVIDADES}}', 'NEE')
    d.save(dst)


if __name__ == '__main__':
    pca_src, micro_src, out_dir = sys.argv[1:4]
    convert_pca(pca_src, f'{out_dir}/pca.docx')
    convert_micro(micro_src, f'{out_dir}/micro.docx')
